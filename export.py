#!/usr/bin/env python3
"""
Export any table to CSV or a formatted .xlsx, into ./exports/.

    python export.py --table poi_overture --format xlsx
    python export.py --table poi_google --format csv --brand Alfamart
    python export.py --all
"""
import argparse
import csv
import os
import sys
import time

import config
import db

TABLES = ["poi_overture", "poi_google", "poi_osm", "poi_unified", "pull_run"]


def stamp():
    return time.strftime("%Y%m%d-%H%M%S")


def fetch(con, table, brand=None, category=None, name=None):
    if table == "pull_run":
        rows = db.runs(con, limit=10000)
    else:
        rows = db.query(con, table, category=category, brand=brand,
                        name_like=name, limit=1_000_000)
    return rows


def to_csv(rows, path):
    if not rows:
        open(path, "w").close()
        return 0
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def to_xlsx(rows, path, sheet):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl is not installed. Run: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    if not rows:
        wb.save(path)
        return 0

    headers = list(rows[0].keys())
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F3864")
    for i, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(h) for h in headers])

    # Width from the widest of the first 200 rows — cheap and good enough.
    for i, h in enumerate(headers, 1):
        widest = max([len(str(h))] +
                     [len(str(r.get(h) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(widest + 2, 48)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--table", choices=TABLES, default="poi_overture")
    ap.add_argument("--all", action="store_true",
                    help="export every table")
    ap.add_argument("--format", choices=["csv", "xlsx"], default="csv")
    ap.add_argument("--brand"), ap.add_argument("--category")
    ap.add_argument("--name")
    args = ap.parse_args()

    if not db.db_exists():
        sys.exit(f"No database at {config.DB_PATH} — run `python migrate.py` "
                 f"then a pull first.")

    con = db.connect()
    tables = TABLES if args.all else [args.table]
    for t in tables:
        rows = fetch(con, t, args.brand, args.category, args.name)
        name = f"{t}-{stamp()}.{args.format}"
        path = os.path.join(config.EXPORT_DIR, name)
        n = (to_xlsx(rows, path, t) if args.format == "xlsx"
             else to_csv(rows, path))
        print(f"  {t:<14}{n:>8,} rows -> exports/{name}")
    con.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
